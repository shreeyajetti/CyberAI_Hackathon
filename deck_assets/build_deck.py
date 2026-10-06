"""Fills the official hackathon template with the Doctor Nexus content.

Slide look stays the template's own (red title bar, Times New Roman text).
Charts, the flow diagram and small visuals use the Modern Minimalist palette
(charcoal / slate / light grey) with the template red as the only accent.
"""
import copy
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "Hackathon PPT Template.pptx"
OUT = ROOT / "Doctor_Nexus_CyberAI_Hackathon.pptx"
SHOTS = ROOT / "Doctor-Nexus-HMS" / "docs"
ASSETS = ROOT / "deck_assets"

# Modern Minimalist + template red
CHARCOAL = RGBColor(0x36, 0x45, 0x4F)
SLATE = RGBColor(0x70, 0x80, 0x90)
LIGHT = RGBColor(0xD3, 0xD3, 0xD3)
TINT = RGBColor(0xF2, 0xF3, 0xF4)
RED = RGBColor(0xC0, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1F, 0x1F, 0x1F)
VIS_FONT = "Arial"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


# ---------- text helpers ----------
def set_runs_text(paragraph, text):
    runs = paragraph.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def fill_body(shape, paras, size=16):
    """Replace a template text box's paragraphs, keeping its first paragraph's styling.
    paras: list of str or dict(t=..., b=bool, color=RGBColor, size=int, gap=pts before)."""
    tf = shape.text_frame
    p0 = tf.paragraphs[0]._p
    r0 = p0.find(A + "r")
    ppr = p0.find(A + "pPr")
    for p in list(tf._txBody.findall(A + "p")):
        tf._txBody.remove(p)
    for i, item in enumerate(paras):
        item = {"t": item} if isinstance(item, str) else item
        p = etree.SubElement(tf._txBody, A + "p")
        if ppr is not None:
            p.append(copy.deepcopy(ppr))
        r = copy.deepcopy(r0) if r0 is not None else etree.SubElement(p, A + "r")
        p.append(r)
        para = tf.paragraphs[-1]
        para.space_before = Pt(0 if i == 0 else item.get("gap", 10))
        run = para.runs[0]
        run.text = item["t"]
        run.font.size = Pt(item.get("size", size))
        run.font.bold = item.get("b", False)
        run.font.color.rgb = item.get("color", INK)


def remove(shape):
    el = shape._element
    el.getparent().remove(el)


def by_name(slide, sid):
    return next(s for s in slide.shapes if s.shape_id == sid)


# ---------- drawing helpers ----------
def box(slide, x, y, w, h, text="", fill=TINT, line=None, size=14, bold=False, color=INK,
        align=PP_ALIGN.LEFT, shape=MSO_SHAPE.RECTANGLE, anchor=MSO_ANCHOR.MIDDLE, font=VIS_FONT,
        margin=0.12):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.08
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    s.shadow.inherit = False
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    lines = text if isinstance(text, list) else [text]
    for i, ln in enumerate(lines):
        ln = {"t": ln} if isinstance(ln, str) else ln
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if i:
            p.space_before = Pt(ln.get("gap", 4))
        r = p.add_run()
        r.text = ln["t"]
        r.font.name = ln.get("font", font)
        r.font.size = Pt(ln.get("size", size))
        r.font.bold = ln.get("b", bold)
        r.font.italic = ln.get("i", False)
        r.font.color.rgb = ln.get("color", color)
    return s


def arrow(slide, x1, y1, x2, y2, color=SLATE, width=1.75, dashed=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if dashed:
        etree.SubElement(ln, A + "prstDash").set("val", "dash")
    etree.SubElement(ln, A + "tailEnd").set("type", "triangle")
    return c


def num_dot(slide, x, y, n, d=0.55, fill=RED):
    return box(slide, x, y, d, d, str(n), fill=fill, size=16, bold=True, color=WHITE,
               align=PP_ALIGN.CENTER, shape=MSO_SHAPE.OVAL, margin=0)


def style_chart_text(chart, size=12):
    chart.font.name = VIS_FONT
    chart.font.size = Pt(size)
    chart.font.color.rgb = CHARCOAL


# ---------- build ----------
prs = Presentation(str(TEMPLATE))

# drop the instructions slide
sld_ids = prs.slides._sldIdLst
first = sld_ids[0]
prs.part.drop_rel(first.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"))
sld_ids.remove(first)

S = list(prs.slides)
notes = {}

# 1. Title
s = S[0]
t = by_name(s, 104).text_frame.paragraphs
set_runs_text(t[0], "Cyber AI Hackathon 2026")
set_runs_text(t[1], "Organized by University of Derby")
title = by_name(s, 108)
title.left, title.width = Inches(0.6), Inches(8.8)
set_runs_text(title.text_frame.paragraphs[0], "Doctor Nexus: An AI Auditor for Missed Surgical Bills")
for row in by_name(s, 106).table.rows:
    if row.cells[0].text.strip() == "Problem Theme":
        set_runs_text(row.cells[2].text_frame.paragraphs[0],
                      "AI in Healthcare & Life Sciences (AI and its applied domain)")
notes[0] = ("Introduce the team and the one-line idea: Doctor Nexus reads surgery notes and the bill, "
            "and tells the hospital what was used but never charged.")

# 2. Background
s = S[1]
fill_body(by_name(s, 115), [
    "During surgery, the team writes down what they used: screws, stents, grafts, clips.",
    "Later, someone in billing types up the bill by hand from those notes.",
    "Things slip through. The note says two stents, the bill says one. Usually nobody spots it until the patient has gone home.",
    {"t": "Who this is for", "b": True, "gap": 18},
    {"t": "Billing teams, hospital admins and finance heads in large (tertiary care) hospitals.", "gap": 4},
])
body = by_name(s, 115)
body.width = Inches(5.2)
box(s, 6.15, 1.25, 3.35, 2.45, [
    {"t": "1–5%", "size": 44, "b": True, "color": RED},
    {"t": "of a hospital's net revenue is lost to charges that never get captured", "size": 13, "color": CHARCOAL},
    {"t": "HFMA estimate", "size": 10, "color": SLATE, "i": True},
], fill=TINT, align=PP_ALIGN.LEFT, margin=0.2)
box(s, 6.15, 3.95, 3.35, 2.45, [
    {"t": "$7.5M", "size": 44, "b": True, "color": CHARCOAL},
    {"t": "recovered in 15 months by one 14-hospital system after fixing its charge capture", "size": 13,
     "color": CHARCOAL},
    {"t": "Novant Health, reported by HFMA", "size": 10, "color": SLATE, "i": True},
], fill=TINT, align=PP_ALIGN.LEFT, margin=0.2)
notes[1] = ("Implants and devices are the most expensive things in an operation, and they are exactly what "
            "gets missed. The numbers on the right are from HFMA (see references). The project README also "
            "quotes 15-20% for Indian hospitals, but we could not find a source for that, so we left it out.")

# 3. Problem statement
s = S[2]
b = by_name(s, 122)
fill_body(b, ["The surgery notes and the bill don't talk to each other. Here's a real record from our test data:"])
b.height = Inches(0.9)
box(s, 0.55, 2.1, 4.0, 0.45, "What the OT note says", fill=CHARCOAL, color=WHITE, bold=True, size=13)
box(s, 0.55, 2.55, 4.0, 1.75, [
    {"t": "Patient P-002 · Cardiac stenting", "size": 11, "color": SLATE},
    {"t": "“2x Onyx Stents (3.5mm) deployed in LAD. 1x Emerge Balloon used for pre-dilation.”",
     "size": 14, "i": True, "gap": 8},
], fill=TINT, anchor=MSO_ANCHOR.TOP, margin=0.18)
box(s, 5.45, 2.1, 4.0, 0.45, "What the bill says", fill=CHARCOAL, color=WHITE, bold=True, size=13)
box(s, 5.45, 2.55, 4.0, 1.75, [
    {"t": "Billing invoice", "size": 11, "color": SLATE},
    {"t": "Onyx Stent  ×  1", "size": 18, "b": True, "gap": 8},
], fill=TINT, anchor=MSO_ANCHOR.TOP, margin=0.18)
arrow(s, 4.62, 3.42, 5.38, 3.42, color=RED, width=2.25)
box(s, 0.55, 4.6, 8.9, 0.75, "Missing from the bill: 1 stent and 1 balloon", fill=None, line=RED,
    color=RED, bold=True, size=18, align=PP_ALIGN.CENTER)
tb = s.shapes.add_textbox(Inches(0.55), Inches(5.6), Inches(8.9), Inches(1.2))
tb.text_frame.word_wrap = True
tb.text_frame.text = "x"
fill_body(tb, [
    "Checking every bill against every note by hand takes hours, so it rarely happens. "
    "Simple keyword checks don't help much either: “Ti screw” and “Titanium Interference Screw” "
    "look like two different items to a computer."], size=15)
for r in tb.text_frame.paragraphs[0].runs:
    r.font.name = "Times New Roman"
notes[2] = "Walk through P-002 slowly. This one example makes the whole problem obvious."

# 4. Objectives
s = S[3]
b = by_name(s, 130)
fill_body(b, ["What we want Doctor Nexus to do:"])
b.height = Inches(0.6)
objectives = [
    ("Read messy OT notes like a person would",
     "Understand brand names, short forms and quantities written in free text."),
    ("Compare every note with its bill",
     "Flag items that were used but not billed, or billed in a smaller quantity."),
    ("Report before the patient is discharged",
     "Give admins a clear recovery table while the bill can still be corrected."),
    ("Keep patient records on the hospital's own machines",
     "Raw files stay local. The AI model only sees what it needs to reason."),
]
for i, (head, sub) in enumerate(objectives):
    y = 1.85 + i * 1.22
    num_dot(s, 0.6, y + 0.12, i + 1)
    box(s, 1.4, y, 8.0, 0.95, [
        {"t": head, "size": 17, "b": True, "font": "Times New Roman"},
        {"t": sub, "size": 14, "color": CHARCOAL, "font": "Times New Roman"},
    ], fill=TINT if i % 2 == 0 else None, anchor=MSO_ANCHOR.MIDDLE, margin=0.15)
notes[3] = "Four objectives. The last one matters a lot for hospitals, because patient data is sensitive."

# 5. Proposed solution
s = S[4]
b = by_name(s, 138)
fill_body(b, [
    "Doctor Nexus is an AI assistant that admins can simply message on Telegram. "
    "Type “Audit this” and it replies with a table of everything that was used in theatre but not charged."])
b.height = Inches(1.3)
steps = [
    ("Read", "Pulls in the OT notes and the bills for the day"),
    ("Match", "Lines up each item in the notes with the bill, even when names differ"),
    ("Flag", "Marks items missing completely or billed short"),
    ("Report", "Sends back a recovery table with patient, item and estimated value"),
]
w, gap, x0, y0 = 2.0, 0.3, 0.55, 2.2
for i, (head, sub) in enumerate(steps):
    x = x0 + i * (w + gap)
    box(s, x, y0, w, 0.65, head, fill=RED if i == 2 else CHARCOAL, color=WHITE, bold=True, size=18,
        align=PP_ALIGN.CENTER)
    box(s, x, y0 + 0.65, w, 1.45, sub, fill=TINT, size=14, color=INK, anchor=MSO_ANCHOR.TOP, margin=0.15)
    if i < 3:
        arrow(s, x + w + 0.03, y0 + 0.33, x + w + gap - 0.03, y0 + 0.33, color=SLATE)
box(s, 0.55, 4.75, 8.9, 1.1, [
    {"t": "Built on the hospital's own data", "b": True, "size": 15},
    {"t": "It works on top of the existing hospital software (HMS). No new system for staff to learn, just a chat.",
     "size": 13, "color": CHARCOAL},
], fill=None, line=LIGHT, margin=0.18)
notes[4] = "Read, match, flag, report. Keep it that simple when presenting."

# 6. Novelty and show stopper
s = S[5]
remove(by_name(s, 146))
rows = [
    ("", "Usual billing checks", "Doctor Nexus"),
    ("Matching", "Exact text only", "Understands short forms and brand names"),
    ("Hidden items", "Only checks what's written", "Knows a hip replacement normally needs bone cement"),
    ("When", "Monthly audit, after discharge", "Same day, before discharge"),
    ("How you ask", "Reports on a desktop", "Plain English, from your phone"),
]
tbl = s.shapes.add_table(len(rows), 3, Inches(0.55), Inches(1.25), Inches(8.9), Inches(3.3)).table
tbl.columns[0].width, tbl.columns[1].width, tbl.columns[2].width = Inches(1.9), Inches(3.2), Inches(3.8)
tbl.first_row = True
for r, row in enumerate(rows):
    for c, val in enumerate(row):
        cell = tbl.cell(r, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = CHARCOAL if r == 0 else (TINT if r % 2 else WHITE)
        cell.margin_left = cell.margin_right = Inches(0.12)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        run = p.add_run()
        run.text = val
        run.font.name = VIS_FONT
        run.font.size = Pt(14)
        run.font.bold = r == 0 or c == 0
        run.font.color.rgb = WHITE if r == 0 else (RED if c == 2 else INK)
box(s, 0.55, 4.9, 8.9, 1.55, [
    {"t": "Show stopper", "size": 13, "b": True, "color": RED},
    {"t": "An admin types “Show me leakage for today's ortho surgeries” on their phone and gets back "
          "a ranked list of missed items, while the bill can still be fixed.", "size": 16, "gap": 6},
], fill=TINT, anchor=MSO_ANCHOR.MIDDLE, margin=0.22)
notes[5] = "The 'implied items' row comes from the project's audit flow document (docs/audi_flow.md)."

# 7. Methodology (flow diagram)
s = S[6]
remove(by_name(s, 153))
# local zone
zone = box(s, 2.75, 1.2, 4.35, 3.95, "", fill=None, line=SLATE)
zone.line.dash_style = 4  # dash
box(s, 2.85, 1.25, 4.1, 0.35, "On the hospital's machine (OpenClaw gateway)", fill=None, size=11,
    color=SLATE, bold=True, margin=0.04)
box(s, 0.45, 2.05, 1.85, 1.1, [{"t": "Admin", "b": True, "size": 14},
                               {"t": "asks on Telegram", "size": 11, "color": CHARCOAL}],
    fill=TINT, align=PP_ALIGN.CENTER, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
box(s, 3.2, 1.8, 3.45, 0.95, [{"t": "Doctor Nexus agent", "b": True, "size": 14, "color": WHITE},
                              {"t": "follows the auditor prompt", "size": 11, "color": WHITE}],
    fill=CHARCOAL, align=PP_ALIGN.CENTER, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
box(s, 3.2, 3.55, 1.6, 1.3, [{"t": "OT notes", "b": True, "size": 13},
                             {"t": "free text", "size": 11, "color": CHARCOAL}],
    fill=TINT, align=PP_ALIGN.CENTER, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
box(s, 5.05, 3.55, 1.6, 1.3, [{"t": "Bills", "b": True, "size": 13},
                              {"t": "structured items", "size": 11, "color": CHARCOAL}],
    fill=TINT, align=PP_ALIGN.CENTER, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
box(s, 7.55, 1.8, 2.0, 0.95, [{"t": "Gemini", "b": True, "size": 14},
                              {"t": "reasoning in the cloud", "size": 11, "color": CHARCOAL}],
    fill=TINT, line=LIGHT, align=PP_ALIGN.CENTER, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
box(s, 0.45, 3.85, 1.85, 1.0, [{"t": "Recovery table", "b": True, "size": 13, "color": WHITE},
                               {"t": "back on the phone", "size": 11, "color": WHITE}],
    fill=RED, align=PP_ALIGN.CENTER, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
arrow(s, 2.32, 2.45, 3.15, 2.3)                 # admin -> agent
arrow(s, 4.0, 3.5, 4.0, 2.8)                    # notes -> agent
arrow(s, 5.85, 3.5, 5.85, 2.8)                  # bills -> agent
arrow(s, 6.7, 2.15, 7.5, 2.15)                  # agent -> gemini
arrow(s, 7.5, 2.5, 6.7, 2.5, dashed=True)       # gemini -> agent
arrow(s, 3.15, 2.6, 2.32, 4.15, color=RED)      # agent -> table
box(s, 4.55, 2.95, 0.9, 0.3, "read_file", fill=WHITE, size=10, color=SLATE, align=PP_ALIGN.CENTER, margin=0)
labels = [
    "Admin sends a message to the Telegram bot.",
    "The agent reads the day's OT notes and bills with its read_file tool.",
    "Gemini compares them: missing items, short quantities, implied items.",
    "Each gap gets a value and a certainty (implants = high).",
    "A formatted recovery table goes back to the admin.",
]
for i, t in enumerate(labels):
    y = 5.3 + i * 0.3
    box(s, 0.55, y, 0.28, 0.28, str(i + 1), fill=CHARCOAL, size=10, bold=True, color=WHITE,
        align=PP_ALIGN.CENTER, shape=MSO_SHAPE.OVAL, margin=0)
    box(s, 0.95, y - 0.03, 8.5, 0.34, t, fill=None, size=13, margin=0.02, font="Times New Roman")
notes[6] = ("Raw files stay inside the dashed box. Only the text needed for reasoning goes out to Gemini "
            "over an encrypted API call. Full architecture diagram is in docs/HLD.jpg in the repo.")

# 8. Tech stack and budget
s = S[7]
fill_body(by_name(s, 160), [
    {"t": "Tech stack used", "b": True, "size": 18},
    {"t": "OpenClaw: runs the agent locally and links it to Telegram", "gap": 12},
    "Google Gemini (2.0 Flash): reads notes and finds gaps",
    "Telegram Bot: the chat screen admins use",
    "JSON files: sample OT notes and bills for 25 patients",
    "Python: our offline test runner (no API key needed)",
], size=15)
fill_body(by_name(s, 163), [
    {"t": "Estimated Budget", "b": True, "size": 18},
    {"t": "Prototype: ₹0", "b": True, "size": 22, "color": RED, "gap": 12},
    {"t": "Everything runs on free tiers and a normal laptop.", "gap": 4},
    {"t": "For a hospital pilot we'd need:", "b": True, "gap": 16},
    "A small on-site server or VM",
    "Paid Gemini API usage (charged per request)",
    "A link to the hospital's billing system (HMS)",
    {"t": "Pilot costs to be confirmed with the partner hospital.", "size": 12, "color": SLATE, "gap": 14},
], size=15)
for sid in (160, 163):
    for p in by_name(s, sid).text_frame.paragraphs:
        p.alignment = PP_ALIGN.LEFT
notes[7] = "Prototype cost is genuinely zero. Don't promise pilot numbers we haven't priced."

# 9. Deliverable / expected outcome
s = S[8]
b = by_name(s, 168)
fill_body(b, [
    "What we deliver: a Telegram chatbot backed by a local AI agent, plus a recovery table admins can act on.",
    {"t": "We ran the audit on the 25 sample patients in the project:", "gap": 8},
], size=15)
b.height = Inches(1.4)
cd = CategoryChartData()
cd.categories = ["Units used in theatre", "Units on the bills"]
cd.add_series("Units", (77, 31))
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.5), Inches(2.6), Inches(5.0), Inches(2.9), cd)
ch = gf.chart
style_chart_text(ch)
ch.has_legend = False
ch.has_title = True
ch.chart_title.text_frame.text = "46 units used but never billed"
tr = ch.chart_title.text_frame.paragraphs[0].runs[0]
tr.font.size, tr.font.bold, tr.font.name, tr.font.color.rgb = Pt(14), True, VIS_FONT, INK
plot = ch.plots[0]
plot.gap_width = 60
plot.has_data_labels = True
plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
plot.data_labels.font.size = Pt(14)
plot.data_labels.font.bold = True
ser = plot.series[0]
for idx, col in enumerate((CHARCOAL, RED)):
    pt = ser.points[idx]
    pt.format.fill.solid()
    pt.format.fill.fore_color.rgb = col
ch.value_axis.visible = False
ch.value_axis.has_major_gridlines = False
ch.value_axis.maximum_scale = 90
ch.category_axis.format.line.color.rgb = LIGHT
ch.category_axis.reverse_order = True

cd2 = CategoryChartData()
cd2.categories = ["Not billed at all", "Billed short"]
cd2.add_series("Gaps", (24, 15))
gf2 = s.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(5.7), Inches(2.6), Inches(3.8), Inches(2.9), cd2)
ch2 = gf2.chart
style_chart_text(ch2)
ch2.has_title = True
ch2.chart_title.text_frame.text = "39 gaps, by type"
tr = ch2.chart_title.text_frame.paragraphs[0].runs[0]
tr.font.size, tr.font.bold, tr.font.name, tr.font.color.rgb = Pt(14), True, VIS_FONT, INK
ch2.has_legend = True
ch2.legend.position = XL_LEGEND_POSITION.BOTTOM
ch2.legend.include_in_layout = False
ch2.legend.font.size = Pt(12)
p2 = ch2.plots[0]
p2.has_data_labels = True
p2.data_labels.font.size = Pt(14)
p2.data_labels.font.bold = True
p2.data_labels.font.color.rgb = WHITE
for idx, col in enumerate((RED, SLATE)):
    pt = p2.series[0].points[idx]
    pt.format.fill.solid()
    pt.format.fill.fore_color.rgb = col
box(s, 0.55, 5.75, 8.9, 0.8, "Every one of the 25 patients had at least one gap. 29 of the 39 were implants or "
    "devices, which are the expensive ones.", fill=TINT, size=14, margin=0.18, font="Times New Roman")
notes[8] = ("Numbers come from running local_run/run_audit.py on workspace/hospital_logs.json. "
            "Rupee values need a hospital price list, which the repo doesn't include.")

# 10. Team role - clear the template's example ticks, team fills in names
s = S[9]
tbl = by_name(s, 175).table
roles = ["Research & problem", "AI agent & prompt", "Data & testing", "Telegram & OpenClaw", "Pitch & docs"]
for c, role in enumerate(roles, start=1):
    set_runs_text(tbl.cell(0, c).text_frame.paragraphs[0], role)
for r in range(1, len(tbl.rows)):
    for c in range(1, len(tbl.columns)):
        for p in tbl.cell(r, c).text_frame.paragraphs:
            for run in p.runs:
                run.text = ""
notes[9] = "Fill in member names and tick the roles each person actually did."

# 11. Screenshot 1 - Telegram
s = S[10]
fill_body(by_name(s, 184), ["The agent answering on Telegram, running through OpenClaw."], size=15)
by_name(s, 184).height = Inches(0.5)
for i, (img, cap) in enumerate([
    ("Screenshot_20260308-021628.png", "“Audit this” returns a summary and the top 5 cases"),
    ("Screenshot_20260308-021745.png", "“Top 3 bills and how to recover” adds next steps"),
]):
    x = 1.55 + i * 3.7
    pic = s.shapes.add_picture(str(SHOTS / img), Inches(x), Inches(1.7), height=Inches(4.45))
    pic.line.color.rgb = LIGHT
    box(s, x - 0.5, 6.2, pic.width / 914400 + 1.0, 0.5, cap, fill=None, size=12, color=CHARCOAL,
        align=PP_ALIGN.CENTER, margin=0)
notes[10] = ("These screenshots are from the project author's own run. The patient IDs shown (P-102, P-123...) "
             "come from an earlier dataset, not the 25 records now in the repo.")

# 12. Screenshot 2 - local run
s = S[11]
fill_body(by_name(s, 191), ["We also ran the same audit on our laptop, offline, with no API key."], size=15)
by_name(s, 191).height = Inches(0.5)
pic = s.shapes.add_picture(str(ASSETS / "local_run.png"), Inches(0.9), Inches(1.7), height=Inches(4.55))
pic.left = int((Inches(10) - pic.width) / 2)
box(s, 0.55, 1.7 + pic.height / 914400 + 0.1, 8.9, 0.4,
    "Rows in red are items that were never billed at all.", fill=None, size=12, color=CHARCOAL, margin=0)
notes[11] = "Command: python run_audit.py (see walkthrough.md). Add --llm to use Gemini instead."

# 13. References
s = S[12]
fill_body(by_name(s, 198), [
    "1.  Doctor Nexus HMS, GitHub repository. github.com/srikanth-hn/Doctor-Nexus-HMS",
    "2.  JTS Health Partners, “How Much Is Your Hospital Losing Through a Poor Charge Capture Process?” "
    "(cites HFMA: 1–5% of net revenue). jtshealthpartners.com/blog",
    "3.  Commure, “Charge Capture in Healthcare: What It Is, Why Revenue Leaks, and How AI Fixes It” "
    "(cites HFMA, Novant Health case). commure.com/blog",
    "4.  OpenClaw documentation. docs.openclaw.ai",
    "5.  Google, Gemini API documentation. ai.google.dev/gemini-api/docs",
    "6.  Telegram, Bot API documentation. core.telegram.org/bots/api",
], size=15)
notes[12] = "Full links are in explanation.md."

# 14. Thank you - keep as is, add a contact line
s = S[13]
box(s, 0.87, 3.55, 8.25, 0.5, "Questions?  github.com/srikanth-hn/Doctor-Nexus-HMS", fill=None, size=14,
    color=SLATE, align=PP_ALIGN.CENTER, font="Times New Roman")

for i, text in notes.items():
    S[i].notes_slide.notes_text_frame.text = text

prs.save(str(OUT))
print("saved", OUT)
